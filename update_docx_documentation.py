"""
Script to format Prototyping Workshop Nicholas Ho.docx with professional,
concise explanations accompanying the 3 chosen screenshots and detailing
the remaining key unpictured states.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text_content, label="WHATSAPP ACTION MESSAGE"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)
    
    # Border: subtle left border
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="111827"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_label = p.add_run(f"[{label}]\n")
    run_label.font.name = "Calibri"
    run_label.font.size = Pt(8.5)
    run_label.font.bold = True
    run_label.font.color.rgb = RGBColor(75, 85, 99)
    
    run_text = p.add_run(text_content)
    run_text.font.name = "Consolas"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(17, 24, 39)
    run_text.font.bold = True

def update_document():
    orig_path = "Prototyping Workshop Nicholas Ho_backup.docx"
    target_path = "Prototyping Workshop Nicholas Ho.docx"

    # Extract images from backup
    orig_doc = docx.Document(orig_path)
    images_data = {}
    for rel in orig_doc.part.rels.values():
        if "image" in rel.target_ref:
            images_data[os.path.basename(rel.target_ref)] = rel.target_part.blob

    # Build fresh document with consistent styling
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)    # Dark slate
    COLOR_MUTED = RGBColor(100, 116, 139)   # Muted gray
    COLOR_TEXT = RGBColor(51, 65, 85)       # Body text

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("SIRCLO Marketplace Monitor")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("State Documentation & System Architecture Walkthrough")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = COLOR_MUTED

    # Project Context & North Star
    h_overview = doc.add_heading(level=2)
    h_overview.paragraph_format.space_before = Pt(10)
    h_overview.paragraph_format.space_after = Pt(4)
    run_ho = h_overview.add_run("1. Executive Summary & North Star")
    run_ho.font.name = "Calibri"
    run_ho.font.size = Pt(13)
    run_ho.font.bold = True
    run_ho.font.color.rgb = COLOR_PRIMARY

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(10)
    p_desc.paragraph_format.line_spacing = 1.15
    run_desc = p_desc.add_run(
        "SIRCLO Marketplace Monitor is a real-time analytics web dashboard built for Brand Account Managers "
        "managing multiple brands across Shopee and TikTok Shop during midnight mega-sale campaigns (e.g., 10.10). "
        "During flash sale rushes, high order volumes and stacked vouchers routinely cause accidental below-cost selling (jual rugi), "
        "while marketplace channel stock runs out even when central warehouses still hold ample inventory. "
        "By replacing the 45-minute manual Google Sheets VLOOKUP process with instant multi-format data ingestion (<2 seconds), "
        "the tool detects profit leaks, forecasts stockout times, and enables 1-click WhatsApp escalation to operations."
    )
    run_desc.font.name = "Calibri"
    run_desc.font.size = Pt(10)
    run_desc.font.color.rgb = COLOR_TEXT

    # -------------------------------------------------------------------------
    # STATE 1: Initial / Empty State
    # -------------------------------------------------------------------------
    h1 = doc.add_heading(level=2)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    run_h1 = h1.add_run("2. State 01: Initial / Empty Dashboard View")
    run_h1.font.name = "Calibri"
    run_h1.font.size = Pt(13)
    run_h1.font.bold = True
    run_h1.font.color.rgb = COLOR_PRIMARY

    # Add Image 3 (Empty State)
    if "image3.png" in images_data:
        temp_img_path = "temp_img3.png"
        with open(temp_img_path, "wb") as f:
            f.write(images_data["image3.png"])
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        run_img = p_img.add_run()
        run_img.add_picture(temp_img_path, width=Inches(6.4))
        os.remove(temp_img_path)

    p1_desc = doc.add_paragraph()
    p1_desc.paragraph_format.space_before = Pt(2)
    p1_desc.paragraph_format.space_after = Pt(12)
    p1_desc.paragraph_format.line_spacing = 1.15
    run_p1 = p1_desc.add_run(
        "How It Works & Key Elements:\n"
        "• Clean Aesthetic & No-Emoji Policy: Uses an uncluttered Stripe/Vercel design system with semantic status dots and tabular typography (Inter).\n"
        "• Multi-Format Dropzones: Four distinct upload slots (Shopee Orders, TikTok Orders, Shopee Inventory, TikTok Inventory) accept CSV, Excel (.xlsx, .xls), and TSV without requiring manual spreadsheet conversion. Numeric cleaners automatically handle Indonesian thousand separators (e.g. '129.000').\n"
        "• Campaign Budget Setup: Allows setting campaign budget parameters (default Rp 5.000.000) for real-time loss tracking.\n"
        "• Instant Prototype Demo: The [Muat Data Sample 10.10] secondary action immediately populates realistic 10.10 campaign transaction data for zero-barrier testing."
    )
    run_p1.font.name = "Calibri"
    run_p1.font.size = Pt(9.5)
    run_p1.font.color.rgb = COLOR_TEXT

    # -------------------------------------------------------------------------
    # STATE 2: Master Data Management (COGS Drawer Open)
    # -------------------------------------------------------------------------
    h2 = doc.add_heading(level=2)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)
    run_h2 = h2.add_run("3. State 02: Master Data Management (COGS / HPP)")
    run_h2.font.name = "Calibri"
    run_h2.font.size = Pt(13)
    run_h2.font.bold = True
    run_h2.font.color.rgb = COLOR_PRIMARY

    # Add Image 2 (COGS Drawer)
    if "image2.png" in images_data:
        temp_img_path = "temp_img2.png"
        with open(temp_img_path, "wb") as f:
            f.write(images_data["image2.png"])
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        run_img = p_img.add_run()
        run_img.add_picture(temp_img_path, width=Inches(6.4))
        os.remove(temp_img_path)

    p2_desc = doc.add_paragraph()
    p2_desc.paragraph_format.space_before = Pt(2)
    p2_desc.paragraph_format.space_after = Pt(12)
    p2_desc.paragraph_format.line_spacing = 1.15
    run_p2 = p2_desc.add_run(
        "How It Works & Key Elements:\n"
        "• Persistent Storage: Master HPP (Harga Pokok Penjualan) data is stored in a local SQLite database (50 registered SKUs), persisting across browser reloads.\n"
        "• In-Place Inline Editing: Operators can adjust the HPP of individual SKUs directly in table input fields and save changes instantly with the [Simpan] button.\n"
        "• Bulk Master Import: Supports uploading replacement master HPP spreadsheets (.csv/.xlsx) to update the entire catalog in bulk.\n"
        "• Decoupled Architecture: Cost master data remains safely separated from high-velocity transaction files, preventing accidental data overrides during midnight rushes."
    )
    run_p2.font.name = "Calibri"
    run_p2.font.size = Pt(9.5)
    run_p2.font.color.rgb = COLOR_TEXT

    # -------------------------------------------------------------------------
    # STATE 3: Analysis Results (Tab 1: Jual Rugi)
    # -------------------------------------------------------------------------
    h3 = doc.add_heading(level=2)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)
    run_h3 = h3.add_run("4. State 03: Live Analysis Results — Margin Loss & Budget Tracking")
    run_h3.font.name = "Calibri"
    run_h3.font.size = Pt(13)
    run_h3.font.bold = True
    run_h3.font.color.rgb = COLOR_PRIMARY

    # Add Image 1 (Results Jual Rugi)
    if "image1.png" in images_data:
        temp_img_path = "temp_img1.png"
        with open(temp_img_path, "wb") as f:
            f.write(images_data["image1.png"])
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        run_img = p_img.add_run()
        run_img.add_picture(temp_img_path, width=Inches(6.4))
        os.remove(temp_img_path)

    p3_desc = doc.add_paragraph()
    p3_desc.paragraph_format.space_before = Pt(2)
    p3_desc.paragraph_format.space_after = Pt(8)
    p3_desc.paragraph_format.line_spacing = 1.15
    run_p3 = p3_desc.add_run(
        "How It Works & Key Elements:\n"
        "• Dynamic Budget Safe Meter: Calculates cumulative voucher losses (sum of (HPP - actual selling price) * quantity) and compares against the campaign budget (Rp 323.221 / Rp 5.000.000, 6.5% used). Displays dynamic color indicators (Aman, Waspada, or Over Budget).\n"
        "• Jual Rugi Alert Cards: Every order where actual selling price is below unit HPP is flagged. Sorted with highest unit loss at the top (e.g. SKU-MOISTURIZER-01 losing Rp 32.658/unit).\n"
        "• Multi-Channel Attribution: Clearly tags whether the loss occurred on Shopee or TikTok Shop, displaying normal price, discounted selling price, unit HPP, and margin delta.\n"
        "• Tab Counts: Dynamically displays detected issues across tabs (100 Jual Rugi, 15 Prediksi Habis, 11 Rebalancing)."
    )
    run_p3.font.name = "Calibri"
    run_p3.font.size = Pt(9.5)
    run_p3.font.color.rgb = COLOR_TEXT

    # -------------------------------------------------------------------------
    # UNPICTURED BUT ESSENTIAL STATES
    # -------------------------------------------------------------------------
    h4 = doc.add_heading(level=2)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(4)
    run_h4 = h4.add_run("5. Essential Workflow States (Unpictured Features)")
    run_h4.font.name = "Calibri"
    run_h4.font.size = Pt(13)
    run_h4.font.bold = True
    run_h4.font.color.rgb = COLOR_PRIMARY

    p4_intro = doc.add_paragraph()
    p4_intro.paragraph_format.space_before = Pt(0)
    p4_intro.paragraph_format.space_after = Pt(6)
    run_p4_intro = p4_intro.add_run(
        "While not pictured in the selected screenshots, the following states and capabilities form critical pillars of the system's operational loop:"
    )
    run_p4_intro.font.name = "Calibri"
    run_p4_intro.font.size = Pt(9.5)
    run_p4_intro.font.italic = True
    run_p4_intro.font.color.rgb = COLOR_MUTED

    # Tab 2: Prediksi Habis
    p_ph = doc.add_paragraph()
    p_ph.paragraph_format.space_before = Pt(4)
    p_ph.paragraph_format.space_after = Pt(4)
    p_ph.paragraph_format.line_spacing = 1.15
    run_ph_title = p_ph.add_run("A. Tab 2 — Prediksi Habis (Run-Rate Velocity & Stockout Forecast)\n")
    run_ph_title.font.name = "Calibri"
    run_ph_title.font.size = Pt(10.5)
    run_ph_title.font.bold = True
    run_ph_title.font.color.rgb = COLOR_PRIMARY
    run_ph = p_ph.add_run(
        "• Calculates hourly sales velocity (Total Units Sold ÷ Elapsed Hours) per SKU and channel.\n"
        "• Predicts exact time until stockout (Remaining Stock ÷ Run Rate). Items running out in <2 hours are tagged 'Kritis', 2-6 hours 'Waspada', and >6 hours 'Aman'.\n"
        "• Zero-Velocity Protection: SKUs with zero sales cleanly display '∞' rather than crashing or showing NaN."
    )
    run_ph.font.name = "Calibri"
    run_ph.font.size = Pt(9.5)
    run_ph.font.color.rgb = COLOR_TEXT

    # Tab 3: Rebalancing
    p_reb = doc.add_paragraph()
    p_reb.paragraph_format.space_before = Pt(6)
    p_reb.paragraph_format.space_after = Pt(4)
    p_reb.paragraph_format.line_spacing = 1.15
    run_reb_title = p_reb.add_run("B. Tab 3 — Multichannel Stock Rebalancing Recommendations\n")
    run_reb_title.font.name = "Calibri"
    run_reb_title.font.size = Pt(10.5)
    run_reb_title.font.bold = True
    run_reb_title.font.color.rgb = COLOR_PRIMARY
    run_reb = p_reb.add_run(
        "• Detects channel stock imbalance: Triggers when marketplace stock is depleted (≤ 5 units) while central warehouse stock is abundant (> 50 units).\n"
        "• Actionable Recommendations: Automatically calculates optimal transfer batches (e.g. 'Rekomendasi: Replenish Shopee (+100 unit)') so operators can shift inventory before listings get suspended."
    )
    run_reb.font.name = "Calibri"
    run_reb.font.size = Pt(9.5)
    run_reb.font.color.rgb = COLOR_TEXT

    # WhatsApp Escalation
    p_wa = doc.add_paragraph()
    p_wa.paragraph_format.space_before = Pt(6)
    p_wa.paragraph_format.space_after = Pt(4)
    p_wa.paragraph_format.line_spacing = 1.15
    run_wa_title = p_wa.add_run("C. 1-Click WhatsApp Escalation Workflow\n")
    run_wa_title.font.name = "Calibri"
    run_wa_title.font.size = Pt(10.5)
    run_wa_title.font.bold = True
    run_wa_title.font.color.rgb = COLOR_PRIMARY
    run_wa = p_wa.add_run(
        "• Per-Item Escalation: Clicking any alert card immediately copies a concise, standardized action alert to the operator's clipboard, complete with SKU, loss per unit, selling price, and HPP, with a 3-second toast feedback ('Pesan aksi disalin'):"
    )
    run_wa.font.name = "Calibri"
    run_wa.font.size = Pt(9.5)
    run_wa.font.color.rgb = COLOR_TEXT

    # Callout Box with exact WA text
    add_callout(
        doc,
        "ALERT: SKU-MOISTURIZER-01 (Ceramide Deep Moisture Cream 50g) jual rugi -32658/unit di shopee. Harga jual Rp 47342, HPP Rp 80000. Segera cek voucher.",
        label="1-CLICK COPIED WHATSAPP ALERT"
    )

    p_wa_summary = doc.add_paragraph()
    p_wa_summary.paragraph_format.space_before = Pt(6)
    p_wa_summary.paragraph_format.space_after = Pt(4)
    p_wa_summary.paragraph_format.line_spacing = 1.15
    run_was = p_wa_summary.add_run(
        "• Executive Summary: The [Salin Ringkasan] button copies a comprehensive broadcast briefing covering overall budget spend, top loss leaders, imminent stockouts, and transfer orders ready for managerial chat groups."
    )
    run_was.font.name = "Calibri"
    run_was.font.size = Pt(9.5)
    run_was.font.color.rgb = COLOR_TEXT

    # Guardrails & Error States
    p_guard = doc.add_paragraph()
    p_guard.paragraph_format.space_before = Pt(6)
    p_guard.paragraph_format.space_after = Pt(4)
    p_guard.paragraph_format.line_spacing = 1.15
    run_guard_title = p_guard.add_run("D. Validation, Guardrails & Error States\n")
    run_guard_title.font.name = "Calibri"
    run_guard_title.font.size = Pt(10.5)
    run_guard_title.font.bold = True
    run_guard_title.font.color.rgb = COLOR_PRIMARY
    run_guard = p_guard.add_run(
        "• Missing Files Guardrail: Prevents partial calculations if fewer than 4 files are uploaded, showing an inline negative error banner.\n"
        "• Schema Validation: Catches corrupt files or missing headers (e.g. missing 'sku_reference_no' in Shopee Orders).\n"
        "• Unmapped SKU Warning: If orders include newly added SKUs missing from the COGS database, a warning banner alerts the operator without aborting calculation for the rest of the catalog.\n"
        "• Rate-Limit Protection (HTTP 429): Enforces a server-side cooldown with live countdown timer ('Tunggu 5s...') to prevent server saturation during multi-operator access."
    )
    run_guard.font.name = "Calibri"
    run_guard.font.size = Pt(9.5)
    run_guard.font.color.rgb = COLOR_TEXT

    doc.save(target_path)
    print("Successfully updated Prototyping Workshop Nicholas Ho.docx!")

if __name__ == "__main__":
    update_document()
